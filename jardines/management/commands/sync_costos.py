import os
import re
import json
import ssl
import urllib.request
from datetime import datetime, date
from django.core.management.base import BaseCommand
from django.conf import settings
from jardines.models import CostoDocente

class Command(BaseCommand):
    help = "Sincroniza los costos laborales docentes desde Google Sheets API / Excel a la base de datos local CostoDocente."

    def add_arguments(self, parser):
        parser.add_argument('--mes', type=str, help='Mes objetivo en formato YYYY-MM (ej: 2026-03)')
        parser.add_argument('--excel', type=str, help='Ruta a archivo Excel local opcional')
        parser.add_argument('--url', type=str, help='URL de Google Apps Script opcional')
        parser.add_argument('--skip-url', action='store_true', help='Omitir consulta a Google Apps Script')

    def handle(self, *args, **options):
        mes_filtro = options.get('mes')
        excel_path_arg = options.get('excel')
        url_arg = options.get('url')
        skip_url = options.get('skip_url')

        registros_creados = 0
        registros_actualizados = 0

        # --- 1. PROCESAR EXCEL LOCAL PRIMERO ---
        excel_candidates = []
        if excel_path_arg:
            excel_candidates.append(excel_path_arg)
        excel_candidates.extend([
            os.path.join(settings.BASE_DIR, "Locaciones 02. Secretaria de Gestión y Participación Ciudadana.xlsx"),
            os.path.join(settings.BASE_DIR, "data", "Locaciones 02. Secretaria de Gestión y Participación Ciudadana.xlsx")
        ])

        excel_path = next((p for p in excel_candidates if os.path.exists(p)), None)
        if excel_path:
            self.stdout.write(self.style.NOTICE(f"Procesando planilla Excel local: {excel_path}"))
            try:
                import openpyxl
                wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
                sheet_names = wb.sheetnames
                sheet_name = 'Locaciones' if 'Locaciones' in sheet_names else sheet_names[0]
                sheet = wb[sheet_name]

                header_row = [cell for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
                
                col_meses = {}
                for idx, val in enumerate(header_row):
                    if hasattr(val, 'year') and hasattr(val, 'month'):
                        col_meses[idx] = f"{val.year:04d}-{val.month:02d}"

                for row in sheet.iter_rows(min_row=2, values_only=True):
                    try:
                        val_dni = row[2] if len(row) > 2 else None
                        if not val_dni:
                            continue

                        dni_clean = self.clean_dni(val_dni)
                        if not dni_clean or len(dni_clean) < 6:
                            continue

                        dnis_to_save = [dni_clean]
                        if len(dni_clean) == 11 and dni_clean[:2] in ('20', '27', '23', '24', '25', '26'):
                            dnis_to_save.append(dni_clean[2:10])

                        for col_idx, m_str in col_meses.items():
                            if mes_filtro and m_str != mes_filtro:
                                continue

                            if col_idx < len(row):
                                c_val = self.parse_costo(row[col_idx])
                                if c_val > 0:
                                    for d_key in dnis_to_save:
                                        obj, created = CostoDocente.objects.update_or_create(
                                            dni=d_key,
                                            mes=m_str,
                                            defaults={
                                                'costo': c_val,
                                                'origen': 'ExcelLocal'
                                            }
                                        )
                                        if created:
                                            registros_creados += 1
                                        else:
                                            registros_actualizados += 1
                    except Exception:
                        pass
                wb.close()
            except Exception as e:
                self.stdout.write(self.style.WARNING(f"Error al procesar Excel local: {e}"))

        # --- 2. PROCESAR GOOGLE SHEETS API (SI NO SE OMITIÓ) ---
        if not skip_url:
            urls_to_process = []
            if url_arg:
                urls_to_process.append((url_arg, "CLI_URL"))
            else:
                url1 = getattr(settings, 'GOOGLE_APPS_SCRIPT_COSTOS_URL', '')
                url2 = getattr(settings, 'GOOGLE_APPS_SCRIPT_COSTOS_URL_2', '')
                if not url1 and not url2:
                    url1 = "https://script.google.com/macros/s/AKfycbxqC9jvLKgZ_rTLraYGkbILfH-en1p0b2Wp1Rl_bl_wLK8_MajcmtBivK28R683ZIwy/exec"
                if url1: urls_to_process.append((url1, "GoogleSheets_Primary"))
                if url2: urls_to_process.append((url2, "GoogleSheets_Secondary"))

            ssl_ctx = ssl.create_default_context()
            ssl_ctx.check_hostname = False
            ssl_ctx.verify_mode = ssl.CERT_NONE

            for api_url, tag in urls_to_process:
                self.stdout.write(self.style.NOTICE(f"Consultando fuente {tag}..."))
                try:
                    req = urllib.request.Request(
                        api_url,
                        headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
                    )
                    with urllib.request.urlopen(req, timeout=15.0, context=ssl_ctx) as resp:
                        if resp.status in (200, 302):
                            payload = json.loads(resp.read().decode('utf-8'))
                            data_rows = []
                            if isinstance(payload, dict) and "data" in payload:
                                data_rows = payload["data"]
                            elif isinstance(payload, list):
                                data_rows = payload

                            self.stdout.write(self.style.SUCCESS(f"Recibidos {len(data_rows)} registros de {tag}"))

                            for item in data_rows:
                                dni_raw = item.get("dni", item.get("DNI", ""))
                                mes_raw = item.get("mes", item.get("Mes", ""))
                                costo_raw = item.get("costo", item.get("costo_total", item.get("CMRC", 0)))

                                dni_clean = self.clean_dni(dni_raw)
                                if not dni_clean or len(dni_clean) < 6:
                                    continue

                                clean_mes = self.parse_mes(mes_raw)
                                if not clean_mes:
                                    continue

                                if mes_filtro and clean_mes != mes_filtro:
                                    continue

                                costo_val = self.parse_costo(costo_raw)
                                if costo_val <= 0:
                                    continue

                                dnis_to_save = [dni_clean]
                                if len(dni_clean) == 11 and dni_clean[:2] in ('20', '27', '23', '24', '25', '26'):
                                    dnis_to_save.append(dni_clean[2:10])

                                for d_key in dnis_to_save:
                                    obj, created = CostoDocente.objects.update_or_create(
                                        dni=d_key,
                                        mes=clean_mes,
                                        defaults={
                                            'costo': costo_val,
                                            'origen': tag
                                        }
                                    )
                                    if created:
                                        registros_creados += 1
                                    else:
                                        registros_actualizados += 1

                except Exception as e:
                    self.stdout.write(self.style.WARNING(f"Error al sincronizar desde {tag}: {e}"))

        total_actuales = CostoDocente.objects.count()
        self.stdout.write(self.style.SUCCESS(
            f"✅ Sincronización finalizada con éxito. Creados: {registros_creados}, Actualizados: {registros_actualizados}. Total registros en BD: {total_actuales}"
        ))

    def clean_dni(self, val):
        if val is None:
            return ""
        try:
            # Manejar floats de openpyxl como 36133163.0 para evitar 361331630
            num_int = int(float(val))
            return str(num_int)
        except (ValueError, TypeError):
            return "".join(filter(str.isdigit, str(val)))

    def parse_mes(self, val):
        if not val:
            return ""
        val_str = str(val).strip().lower()
        if len(val_str) == 7 and val_str[4] == '-':
            return val_str

        meses_dict = {
            'ene': '01', 'feb': '02', 'mar': '03', 'abr': '04', 'may': '05', 'jun': '06',
            'jul': '07', 'ago': '08', 'sep': '09', 'set': '09', 'oct': '10', 'nov': '11', 'dic': '12'
        }
        for key, num_str in meses_dict.items():
            if key in val_str:
                digits = re.findall(r'\d+', val_str)
                if digits:
                    yr = digits[-1]
                    if len(yr) == 2:
                        yr = "20" + yr
                    return f"{yr}-{num_str}"

        try:
            dt = datetime.strptime(val_str, "%Y-%m-%d")
            return dt.strftime("%Y-%m")
        except ValueError:
            pass
        return val_str

    def parse_costo(self, val):
        if val is None:
            return 0
        if isinstance(val, (int, float)):
            num = float(val)
        else:
            str_val = str(val).strip()
            if ',' in str_val:
                str_val = str_val.replace('.', '').replace(',', '.')
            clean_str = "".join([c for c in str_val if c.isdigit() or c == '.'])
            try:
                num = float(clean_str)
            except ValueError:
                num = 0

        while 0 < num < 100000:
            num *= 1000  # Ajustar montos dados en miles
        return num
